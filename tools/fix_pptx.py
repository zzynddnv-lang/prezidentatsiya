"""Taqdimotdagi kamchiliklarni tuzatuvchi skript.

Bir martalik skript: u Gamma'dan eksport qilingan ASL nusxaga (commit 0e6eeab)
qo'llangan va natija allaqachon repozitoriyada. Qayta ishlatish uchun avval asl
faylni tiklang, so'ng skriptni ishga tushiring (fayl joyida yangilanadi):
    git checkout 0e6eeab -- Biznes-va-tadbirkorlik-Business-Model-Canvas-va-BPM.pptx
    python tools/fix_pptx.py
Talab: pip install python-pptx
"""
import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "Biznes-va-tadbirkorlik-Business-Model-Canvas-va-BPM.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def replace_text(slide, old, new):
    hits = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    hits += 1
    assert hits, f"topilmadi: {old!r}"


def set_text(shape, text):
    runs = shape.text_frame.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def clone(slide, shape):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def normalize_apostrophes(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        t = r.text.replace("ʻ", "'").replace("ʼ", "'").replace("’", "'")
                        if t != r.text:
                            r.text = t


def remove_gamma_badge(prs):
    for layout in prs.slide_layouts:
        for sh in list(layout.shapes):
            if sh._element.xpath(".//a:hlinkClick"):
                sh._element.getparent().remove(sh._element)


def fix_slide6(slide):
    """O'ng tomondagi ro'yxatni 5 tadan 9 ta blokga to'ldirish."""
    shapes = list(slide.shapes)
    row = shapes[3:6]  # chiziq, sarlavha, tavsif (1-qator); raqam sarlavhaga qo'shiladi
    for sh in shapes[2:3] + shapes[6:]:
        sh._element.getparent().remove(sh._element)
    blocks = [
        ("Customer Segments", "Mijozlaringiz kim?"),
        ("Value Propositions", "Ularga qanday qiymat berasiz?"),
        ("Channels", "Mijozlarga qanday yetib borasiz?"),
        ("Customer Relationships", "Mijozlar bilan qanday munosabat quriladi?"),
        ("Revenue Streams", "Daromad qayerdan keladi?"),
        ("Key Resources", "Qanday asosiy resurslar kerak?"),
        ("Key Activities", "Qanday asosiy ishlar bajariladi?"),
        ("Key Partnerships", "Asosiy hamkorlar kimlar?"),
        ("Cost Structure", "Asosiy xarajatlar qanday?"),
    ]
    base = [sh.top for sh in row]
    top0, stride = Inches(1.25), Inches(0.62)
    for i, (title, desc) in enumerate(blocks):
        items = row if i == 0 else [clone(slide, sh) for sh in row]
        for sh, b in zip(items, base):
            sh.top = Emu(top0 + stride * i + (b - base[0]))
        set_text(items[1], f"{i + 1:02d}  {title}")
        set_text(items[2], desc)


def fix_slide7(slide):
    """Onlayn kurs misolini 6 tadan 9 ta BMC blokiga kengaytirish (3x3 to'r)."""
    shapes = list(slide.shapes)
    card = shapes[2:5]  # ramka, sarlavha, tavsif
    cards_src = [shapes[2 + 3 * k: 5 + 3 * k] for k in range(6)]
    # 4-karta (Daromad) ramkasi xira bo'lib qolgan — boshqalariga o'xshatamiz
    good_ln = copy.deepcopy(cards_src[0][0]._element.spPr.find(A + "ln"))
    for sh in shapes[2:]:
        sh._element.getparent().remove(sh._element)
    data = [
        ("Mijozlar", "O'quvchilar, talabalar, kattalar"),
        ("Qiymat", "Sifatli ta'lim, qulay vaqtda o'qish"),
        ("Kanallar", "Veb-sayt, mobil ilova, YouTube"),
        ("Munosabatlar", "Mentor yordami, chat, jamoa"),
        ("Daromad", "Kurs to'lovi, obuna, sertifikat"),
        ("Resurslar", "O'qituvchilar, server, kontent"),
        ("Asosiy faoliyat", "Dars tayyorlash, marketing"),
        ("Hamkorlar", "Universitetlar, IT-kompaniyalar"),
        ("Xarajatlar", "Server, maosh, reklama"),
    ]
    left0, top0 = Inches(0.79), Inches(2.3)
    gap = Inches(0.2)
    w = (Inches(11.76) - 2 * gap) // 3
    h = Inches(1.4)
    dx_title, dy_title = card[1].left - card[0].left, card[1].top - card[0].top
    dy_desc = Inches(0.62)
    for i, (title, desc) in enumerate(data):
        src = cards_src[i % 6]
        if i % 6 == 3:
            src = cards_src[0]
        frame, t, d = [clone(slide, sh) for sh in src]
        ln = frame._element.spPr.find(A + "ln")
        if ln is not None and i % 6 == 3:
            frame._element.spPr.replace(ln, copy.deepcopy(good_ln))
        x = left0 + (w + gap) * (i % 3)
        y = top0 + (h + gap) * (i // 3)
        frame.left, frame.top, frame.width, frame.height = Emu(x), Emu(y), Emu(w), Emu(h)
        for sh in (t, d):
            sh.left, sh.width = Emu(x + dx_title), Emu(w - 2 * dx_title)
            sh.text_frame.word_wrap = True
        t.top = Emu(y + dy_title)
        d.top = Emu(y + dy_desc)
        d.height = Emu(Inches(0.65))
        set_text(t, title)
        set_text(d, desc)
        d.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    replace_text(slide, "ko'rib chiqamiz:", "9 ta blok bo'yicha ko'rib chiqamiz:")


def main():
    prs = Presentation(DECK)
    s = prs.slides
    remove_gamma_badge(prs)
    normalize_apostrophes(prs)

    replace_text(s[1], "Dukonda sotish", "Do'konda sotish")
    replace_text(s[2], "saloon", "salon")
    replace_text(s[3], "Ma'lumotlangan xaridorlar", "Mahsulotni sotib oluvchilar")
    fix_slide6(s[5])
    fix_slide7(s[6])
    replace_text(s[8], "To'lov maydoni nazorat qilinadi", "To'lov holati tasdiqlanadi")

    src = s[9].shapes[-1]
    p = src.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = "Foydalanilgan manbalar: "
    runs[1].text = (
        "A. Osterwalder, Y. Pigneur — «Business Model Generation» (2010), Strategyzer  |  "
        "BPM — ABPMP BPM CBOK, WfMC, IBM  |  Tadbirkorlik — O'zbekiston Respublikasining "
        "«Tadbirkorlik faoliyati erkinligining kafolatlari to'g'risida»gi Qonuni"
    )
    for r in runs[2:]:
        r._r.getparent().remove(r._r)
    src.height = Emu(Inches(0.55))

    prs.save(DECK)
    print("Saqlandi:", DECK.name)


if __name__ == "__main__":
    main()
